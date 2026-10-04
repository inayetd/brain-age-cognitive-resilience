"""The H1 results slide, plus a backup slide for the detail that does not
belong on it.

Layout: figure on the left, three findings on the right, one footnote at
the bottom. The findings are the figure in words, so that a reader who
looks at the text and a reader who looks at the plot arrive at the same
place; nothing is said on the right that cannot be seen on the left.

Model fit and factor loadings are deliberately not here. They go on the
backup slide, where they are available if asked and out of the way if not.

Typography follows the other decks in presentations/.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

FIGD = sys.argv[1]
OUT = sys.argv[2]

NAVY = RGBColor(0x16, 0x32, 0x4F)
GREY = RGBColor(0x5A, 0x65, 0x70)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x15, 0x49, 0x7B)
PURPLE = RGBColor(0x7B, 0x2D, 0x8E)
MINUS = "−"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def textbox(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, parts, size, color, font="Calibri", bold=False, first=False,
         space_before=0, align=PP_ALIGN.LEFT, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    if isinstance(parts, str):
        parts = [(parts, False)]
    for part in parts:
        t, it = part[0], part[1]
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.color.rgb = part[2] if len(part) > 2 else color
        r.font.name = font
        r.font.italic = it
        r.font.bold = bold
    return p


def picture(s, path, left, top, max_w, max_h):
    iw, ih = Image.open(path).size
    k = min(max_w / iw, max_h / ih)
    w, h = iw * k, ih * k
    s.shapes.add_picture(path, Inches(left + (max_w - w) / 2),
                         Inches(top + (max_h - h) / 2), Inches(w), Inches(h))


# ==================================================================== slide 1
s = prs.slides.add_slide(BLANK)

tf = textbox(s, 0.62, 0.45, 12.1, 0.9)
para(tf, "Higher brain-age gap was associated with lower general cognition",
     27, NAVY, font="Cambria", bold=True, first=True)

fig = os.path.join(FIGD, "fig10_h1_presentation.png")
if not os.path.exists(fig):
    raise SystemExit("missing " + fig)
picture(s, fig, 0.45, 1.45, 7.45, 4.6)

# Right column: the same three facts the figure shows.
tf = textbox(s, 8.15, 1.75, 4.65, 4.2)

para(tf, "Structural BAG", 16, BLUE, bold=True, first=True)
para(tf, [("β = " + MINUS + ".159", False),
          ("   97.5% CI [" + MINUS + ".217, " + MINUS + ".101]", False)],
     15, INK, space_before=2)
para(tf, [("p", True), (" < .001", False)], 15, INK, space_before=1)

para(tf, "Functional (PCA) BAG", 16, PURPLE, bold=True, space_before=16)
para(tf, [("β = " + MINUS + ".065", False),
          ("   97.5% CI [" + MINUS + ".126, " + MINUS + ".004]", False)],
     15, INK, space_before=2)
para(tf, [("p", True), (" = .017", False)], 15, INK, space_before=1)

para(tf, "Both modalities", 16, NAVY, bold=True, space_before=16)
para(tf, [("Negative, as predicted, and both below the "
           "Bonferroni-adjusted threshold ", False),
          ("α", True), (" = .025", False)],
     15, INK, space_before=2, line=1.15)

tf = textbox(s, 0.62, 6.35, 12.1, 0.7)
para(tf, [("Multiple-imputation bifactor S–1 SEM over 50 imputed "
           "datasets, pooled by Rubin’s rules; adjusted for "
           "chronological age, sex and education; ", False),
          ("n", True), (" = 1,342 per modality.", False)],
     12.5, GREY, first=True, line=1.15)

# ==================================================================== slide 2
s = prs.slides.add_slide(BLANK)

tf = textbox(s, 0.62, 0.45, 12.1, 0.95)
para(tf, "H1 — supporting detail", 26, NAVY, font="Cambria",
     bold=True, first=True)
para(tf, "Backup: what the effect is worth in years, and which part of "
         "cognition it tracks", 13.5, GREY, space_before=5)

fig2 = os.path.join(FIGD, "fig11_h1_supporting_presentation.png")
if os.path.exists(fig2):
    picture(s, fig2, 0.45, 1.70, 12.45, 4.35)

tf = textbox(s, 0.62, 6.30, 12.1, 0.9)
para(tf, [("The gap and chronological age are both in years, so the ratio "
           "of their per-year effects is unitless: one year of structural "
           "gap is worth about half a year of ordinary ageing. The "
           "conversion treats the two standard deviations as known, which "
           "understates the uncertainty slightly. Panel b uses the "
           "bifactor structure — a single-factor model could not "
           "separate general from domain-specific variance.", False)],
     12.5, GREY, first=True, line=1.15)

prs.save(OUT)
print("wrote " + OUT)
